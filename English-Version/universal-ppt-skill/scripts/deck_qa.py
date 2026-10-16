#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""deck_qa.py — universal-ppt-skill single-entry QA gate (Stage 5 executable gates)

Turns the prose rules of references/qa-gates.md into machine-decidable checks (G1-G15).
Any FAIL -> exit code = 1.

Bilingual decks: gates G2 (speaker notes), G9 (banned template phrases), G13 (action
titles) and the evidence-length checks auto-detect the language of each note/title.
Chinese decks are measured in characters (default 80-120); English decks in words
(default 60-100). Thresholds are configurable via --notes-min/--notes-max (CN) and
--notes-min-en/--notes-max-en (EN).

Usage:
  python scripts/deck_qa.py <deck.pptx> [--spec work/deck-spec.json]
        [--figures work/figures.json] [--paper-figures work/paper-figures.json]
        [--out work/qa_output.txt]
        [--notes-min 80] [--notes-max 120]
        [--notes-min-en 60] [--notes-max-en 100]
        [--mode A|B|C|D|E|F]
"""

import argparse
import hashlib
import json
import math
import os
import re
import sys

# ---------------------------------------------------------------- dependencies
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn
except ImportError:
    sys.stderr.write(
        "[FATAL] python-pptx is not installed; deck_qa.py cannot run.\n"
        "        Install it first:  pip install python-pptx\n"
    )
    raise SystemExit(2)

# ---------------------------------------------------------------- constants
CANVAS_W_IN_DEFAULT = 13.333
CANVAS_H_IN_DEFAULT = 7.5
BOUND_TOL_IN = 0.03          # shape out-of-bounds tolerance (inches)
GIANT_H_IN = 8.0             # giant-shape threshold
GIANT_W_IN = 13.5
EMU_PER_PT = 12700
EMU_PER_IN = 914400

VISUAL_PAGE_TYPES = {"content", "data", "table"}   # G7 requires a visual anchor (timeline pages are built from node+connector shapes, hence exempt)
EXEMPT_PAGE_TYPES = {"title", "statement", "closing", "section", "timeline"}  # G7 exemptions
TITLE_GATE_TYPES = {"content", "data", "timeline", "closing"}   # G13 scope

PLACEHOLDER_PATTERNS = [
    "XXX", "lorem", "TODO", "TBD", "[insert", "{{", "}}", "placeholder",
    "待补充", "此处", "占位",
]
FORBIDDEN_PATTERNS = [
    # --- Chinese template phrases (G9, CN mode) ---
    r"一句话总结", r"最有价值的后续方向", r"值得注意的是", r"总的来说",
    r"综上所述", r"从某种意义上", r"提供了新的视角", r"具有重要意义",
    r"未来可以进一步探索", r"不是.{0,15}?而是", r"不只是.{0,15}?更是",
    r"众所周知", r"极大地", r"显著提升",
    # --- English template phrases (G9, EN mode) ---
    r"\bin (one )?sentence\b",
    r"\bit is worth noting that\b",
    r"\bit is important to note that\b",
    r"\bit should be noted that\b",
    r"\bin summary\b",
    r"\bin conclusion\b",
    r"\bin a nutshell\b",
    r"\bas mentioned (above|earlier)\b",
    r"\bit goes without saying\b",
    r"\bat the end of the day\b",
    r"\bdelve(s|d)? into\b",
    r"\bin today's fast-paced\b",
    r"\bunlock(ing)? the (full )?potential\b",
    r"\bgame[- ]changer\b",
    r"\bseamless(ly)? integrat(e|es|ing|ion)\b",
    r"\brevolutioniz(e|es|ing|ed)\b",
    r"\bnot only\b.{0,60}?\bbut also\b",
]
TITLE_BLACKLIST = {
    # --- Chinese section-name titles ---
    "研究背景", "背景", "方法", "方法介绍", "方案", "实验结果", "结果", "结论",
    "总结", "创新点", "不足", "参考文献", "致谢", "目录", "汇报大纲", "项目介绍",
    "实验设置",
    # --- English section-name titles ---
    "Background", "Introduction", "Related Work", "Literature Review",
    "Method", "Methods", "Methodology", "Approach", "Proposed Method",
    "Experimental Setup", "Experiments", "Results", "Experimental Results",
    "Conclusion", "Conclusions", "Summary", "Contributions", "Limitations",
    "References", "Acknowledgments", "Acknowledgements", "Outline", "Agenda",
    "Overview", "Thank You", "Thanks", "Q&A", "Questions",
}
CITATION_PATTERNS = [
    r"来源[:：]", r"[Ss]ource[:：]", r"Figure\s*\d+", r"Fig\.\s*\d+",
    r"Table\s*\d+", r"图\s*\d+", r"表\s*\d+",
    r"自绘", r"自制", r"作者绘制",
    r"[Ss]elf[- ]made", r"[Ss]elf[- ]drawn", r"[Dd]rawn by (the )?author",
]

# Figure source types (G10): paper decks must use original figures first; self-made
# figures are only allowed as a supplement, with a documented reason on record.
PAPER_CROP = "pdf-crop"
SELF_MADE = "self-made"
MODE_A_SOURCE_TYPES = {PAPER_CROP, "docs-asset", SELF_MADE}

# G4 line-break points: CJK and ASCII punctuation are legal break points; when looking
# for "unbreakable tokens" they must be split off first.
BREAK_CHARS = "，。、；：？！）】》」』％…—·,.;:?!)]}"

# Language auto-detection (G2/G12/G13/G15 evidence checks)
CJK_CHAR_RE = re.compile(
    r"[\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef]")
EN_WORD_RE = re.compile(
    r"[A-Za-z0-9\u00c0-\u024f]+(?:['\u2019-][A-Za-z0-9\u00c0-\u024f]+)*")


def cjk_char_count(text):
    return len(CJK_CHAR_RE.findall(text or ""))


def is_cjk_text(text, min_cjk=5, ratio=0.3):
    """True = the text is predominantly CJK (apply the CN-mode thresholds)."""
    t = (text or "").strip()
    n = cjk_char_count(t)
    return n >= min_cjk and n >= ratio * max(1, len(t))


def count_english_words(text):
    return len(EN_WORD_RE.findall(text or ""))


def min_evidence_len(text):
    """Minimum length for free-text evidence: >=8 CJK chars, >=20 chars otherwise."""
    return 8 if cjk_char_count(text or "") >= 5 else 20


# ---------------------------------------------------------------- utilities
def walk_shapes(shapes):
    """Recursively expand group shapes, yielding leaf shapes."""
    for sh in shapes:
        try:
            is_group = sh.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:
            is_group = False
        if is_group:
            yield from walk_shapes(sh.shapes)
        else:
            yield sh


def is_picture_shape(sh):
    try:
        return sh.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE)
    except Exception:
        return False


def is_visual_shape(sh):
    """Visual-anchor definition (core of G5/G7): PICTURE / native chart / table.
    Everything else (text boxes, colored blocks, rounded cards, shape-assembled
    'diagrams') does not count."""
    if is_picture_shape(sh):
        return True
    if getattr(sh, "has_chart", False):
        return True
    if getattr(sh, "has_table", False):
        return True
    return False


def shape_bounds_in(sh):
    """Return (left, top, right, bottom) in inches, or None if coordinates are missing."""
    left, top = sh.left, sh.top
    w, h = sh.width, sh.height
    if left is None or top is None or w is None or h is None:
        return None
    return (left / EMU_PER_IN, top / EMU_PER_IN,
            (left + w) / EMU_PER_IN, (top + h) / EMU_PER_IN)


def slide_text(slide):
    parts = []
    for sh in walk_shapes(slide.shapes):
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False):
            for row in sh.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return "\n".join(p for p in parts if p)


def notes_text(slide):
    try:
        if not slide.has_notes_slide:
            return ""
        return slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        return ""


def char_width_em(ch):
    """CJK / full-width characters count 1.0 em, ASCII 0.5 em."""
    return 1.0 if ord(ch) > 0x2E80 else 0.5


def resolve_font_pt(run, para):
    size = getattr(run.font, "size", None) or getattr(para.font, "size", None)
    return size.pt if size is not None else 18.0


def estimate_text_block(sh):
    """Estimate a text frame. Returns (est_h_pt, max_line_w_pt, usable_w_pt, frame_h_pt, frame_w_pt)."""
    tf = sh.text_frame
    frame_w_pt = sh.width / EMU_PER_PT
    frame_h_pt = sh.height / EMU_PER_PT
    ml = tf.margin_left if tf.margin_left is not None else 0
    mr = tf.margin_right if tf.margin_right is not None else 0
    mt = tf.margin_top if tf.margin_top is not None else 0
    mb = tf.margin_bottom if tf.margin_bottom is not None else 0
    usable_w_pt = max(frame_w_pt - (ml + mr) / EMU_PER_PT, 1.0)

    total_lines = 0
    max_line_w = 0.0
    max_font = 0.0
    for para in tf.paragraphs:
        runs = list(para.runs)
        if not runs:
            if not para.text:
                total_lines += 0          # empty paragraphs don't count as lines
                continue
            runs = None
        if runs is None:
            seg_pt = resolve_font_pt(para, para)
            w = sum(char_width_em(c) * seg_pt for c in para.text)
            max_font = max(max_font, seg_pt)
        else:
            w = 0.0
            for r in runs:
                seg_pt = resolve_font_pt(r, para)
                w += sum(char_width_em(c) * seg_pt for c in r.text)
                max_font = max(max_font, seg_pt)
        max_line_w = max(max_line_w, w)
        n_lines = max(1, int(math.ceil(w / usable_w_pt)))
        total_lines += n_lines
    if max_font <= 0:
        max_font = 18.0
    line_h = max_font * 1.35
    est_h_pt = total_lines * line_h + (mt + mb) / EMU_PER_PT
    return est_h_pt, max_line_w, usable_w_pt, frame_h_pt, frame_w_pt


def longest_unbreakable_token(sh):
    """Return (token, width_pt): the longest token that cannot be broken mid-word.
    CJK text can break at any character, so only contiguous ASCII runs of >= 6
    characters (e.g. langgraph, softmax(QK^T/sqrt(d))*V, paths/URLs) count as
    unbreakable tokens — those are the ones that actually shoot out of the frame.
    Note: CJK punctuation (、，。；：) is itself a legal break point and must be
    split off first, otherwise "Autoformer, FEDformer, PatchTST" would be
    misjudged as one giant token."""
    best_tok, best_w = "", 0.0
    splitter = re.compile(r"[\s" + re.escape(BREAK_CHARS) + r"]+")
    for para in sh.text_frame.paragraphs:
        pt = 0.0
        for r in para.runs:
            pt = max(pt, resolve_font_pt(r, para))
        if pt <= 0:
            pt = resolve_font_pt(para, para)
        for tok in splitter.split(para.text):
            ascii_n = sum(1 for c in tok if ord(c) < 128 and not c.isspace())
            if ascii_n < 6:
                continue
            w = sum(char_width_em(c) * pt for c in tok)
            if w > best_w:
                best_tok, best_w = tok, w
    return best_tok, best_w


def max_font_pt_of_shape(sh):
    best = 0.0
    if not sh.has_text_frame:
        return best
    for para in sh.text_frame.paragraphs:
        if para.font.size is not None:
            best = max(best, para.font.size.pt)
        for r in para.runs:
            if r.font.size is not None:
                best = max(best, r.font.size.pt)
    return best


# ---------------------------------------------------------------- main flow
class QA:
    def __init__(self, args):
        self.args = args
        self.lines = []
        self.fail_count = 0
        self.warn_count = 0
        self.used_pdfcrop_labels = set()   # paper figures placed on slides (by Figure label)
        self.used_selfmade = 0             # number of self-made figures placed on slides

    def emit(self, tag, level, msg):
        self.lines.append(f"[{level}] {tag} {msg}")
        if level == "FAIL":
            self.fail_count += 1
        elif level == "WARN":
            self.warn_count += 1

    def run(self, prs, spec, figures, inventory, mode):
        slides = list(prs.slides)
        n = len(slides)
        total_pages = spec.get("pages") if spec else None

        # basic derived data
        per_slide_shapes = [list(walk_shapes(s.shapes)) for s in slides]
        per_slide_visual = [any(is_visual_shape(sh) for sh in shs) for shs in per_slide_shapes]

        self.g1_pages(n, total_pages)
        self.g2_notes(slides)
        self.g3_bounds(slides, per_slide_shapes)
        self.g4_overflow(slides, per_slide_shapes)
        self.g5_visual_ratio(n, per_slide_visual, spec)
        self.g6_consecutive(n, per_slide_visual)
        self.g7_anchor(n, per_slide_visual, spec)
        self.g8_placeholder(slides)
        self.g9_forbidden(slides)
        self.g10_images(slides, figures, spec, mode)
        self.g11_citation(slides, per_slide_shapes, figures, mode)
        self.g12_confirm(spec)
        self.g13_action_title(slides, per_slide_shapes, spec)
        self.g14_fonts_shadow(slides, per_slide_shapes)
        self.g15_paper_figs(inventory, spec, mode)

        self.evidence(n, per_slide_visual, spec)

    # ---- G1 page count ----
    def g1_pages(self, n, total_pages):
        if total_pages is None:
            self.emit("G1", "WARN", f"spec has no 'pages'; skipping page-count check (actual {n} pages)")
        elif n == total_pages:
            self.emit("G1", "PASS", f"slides={n} matches spec.pages")
        else:
            self.emit("G1", "FAIL", f"slides={n} does not match spec.pages={total_pages}")

    # ---- G2 speaker notes length (language-adaptive) ----
    def g2_notes(self, slides):
        lo_cn, hi_cn = self.args.notes_min, self.args.notes_max
        lo_en, hi_en = self.args.notes_min_en, self.args.notes_max_en
        bad = []
        for i, s in enumerate(slides, 1):
            txt = notes_text(s)
            if is_cjk_text(txt):
                n_len, unit, lo, hi = len(txt), "chars(CN)", lo_cn, hi_cn
            else:
                n_len, unit, lo, hi = count_english_words(txt), "words(EN)", lo_en, hi_en
            self.emit("G2", "PASS", f"page {i:02d} notes={n_len} {unit} [range {lo}-{hi}]")
            if n_len < lo or n_len > hi:
                bad.append((i, n_len, unit))
        if bad:
            self.emit("G2", "FAIL",
                      f"notes length out of range (CN [{lo_cn},{hi_cn}] chars / EN [{lo_en},{hi_en}] words): {bad}")
        else:
            self.emit("G2", "PASS",
                      f"all notes within range (CN [{lo_cn},{hi_cn}] chars / EN [{lo_en},{hi_en}] words)")

    # ---- G3 shape out-of-bounds ----
    def g3_bounds(self, slides, per_slide_shapes):
        cw, ch = self.canvas
        bad = []
        giants = []
        for i, shs in enumerate(per_slide_shapes, 1):
            for k, sh in enumerate(shs):
                b = shape_bounds_in(sh)
                if b is None:
                    continue
                l, t, r, bo = b
                if (l < -BOUND_TOL_IN or t < -BOUND_TOL_IN or
                        r > cw + BOUND_TOL_IN or bo > ch + BOUND_TOL_IN):
                    bad.append((i, k, round(l, 2), round(t, 2), round(r, 2), round(bo, 2)))
                w_in = r - l
                h_in = bo - t
                if h_in > GIANT_H_IN or w_in > GIANT_W_IN:
                    giants.append((i, k, round(w_in, 2), round(h_in, 2)))
        if bad:
            self.emit("G3", "FAIL",
                      f"shapes out of bounds (canvas {cw}x{ch}in, tolerance {BOUND_TOL_IN}in): "
                      f"(page,idx,left,top,right,bottom)={bad}")
        else:
            self.emit("G3", "PASS", "no shapes out of bounds")
        if giants:
            self.emit("G3", "WARN",
                      f"GIANT shapes (page,idx,w_in,h_in), threshold h>{GIANT_H_IN} or w>{GIANT_W_IN}: {giants}")
        else:
            self.emit("G3", "PASS", "no GIANT shapes")

    # ---- G4 text overflow estimate ----
    def g4_overflow(self, slides, per_slide_shapes):
        over_h, over_w, dense_para = [], [], []
        for i, shs in enumerate(per_slide_shapes, 1):
            for k, sh in enumerate(shs):
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue
                if sh.width is None or sh.height is None:
                    continue
                est_h, max_line_w, usable_w, frame_h, frame_w = estimate_text_block(sh)
                if est_h > frame_h * 1.15:
                    over_h.append((i, k, round(est_h, 1), round(frame_h, 1)))
                # width is only judged on 'unbreakable long tokens' (CJK can wrap anywhere;
                # judging whole multi-line paragraphs on width would produce false positives)
                tok, tok_w = longest_unbreakable_token(sh)
                if tok_w > usable_w * 1.02:
                    over_w.append((i, k, tok, round(tok_w, 1), round(usable_w, 1)))
                # a paragraph far wider than its container means too many wrapped lines: suggest splitting
                if max_line_w > usable_w * 2.5:
                    dense_para.append((i, k, round(max_line_w, 1), round(usable_w, 1)))
        if over_h:
            self.emit("G4", "FAIL",
                      f"height overflow (page,idx,est_h_pt,frame_h_pt): {over_h}")
        else:
            self.emit("G4", "PASS", "no height overflow")
        if over_w:
            self.emit("G4", "FAIL",
                      f"unbreakable long token exceeds frame (page,idx,token,token_w_pt,usable_w_pt): {over_w}")
        else:
            self.emit("G4", "PASS", "no unbreakable long token exceeding the frame")
        if dense_para:
            self.emit("G4", "WARN",
                      f"paragraph far wider than container (suggest splitting or widening) "
                      f"(page,idx,para_w_pt,usable_w_pt): {dense_para}")
        else:
            self.emit("G4", "PASS", "no overlong single paragraph")

    # ---- G5 visual-page ratio ----
    def g5_visual_ratio(self, n, vis, spec):
        need = 0.6
        if spec and spec.get("visual_ratio_min") is not None:
            need = float(spec["visual_ratio_min"])
        k = sum(1 for v in vis if v)
        ratio = k / n if n else 0.0
        msg = f"visual_pages={k}/{n} = {ratio * 100:.1f}% (need >= {need * 100:.0f}%)"
        if ratio >= need:
            self.emit("G5", "PASS", msg)
        else:
            self.emit("G5", "FAIL", msg + " —— visual pages insufficient (only pictures / original figures / native charts / tables count)")

    # ---- G6 consecutive non-visual pages ----
    def g6_consecutive(self, n, vis):
        runs = []
        i = 0
        while i < n:
            if not vis[i]:
                j = i
                while j < n and not vis[j]:
                    j += 1
                if j - i >= 3:
                    runs.append((i + 1, j))
                i = j
            else:
                i += 1
        if runs:
            self.emit("G6", "FAIL", f"consecutive non-visual pages >= 3: longest runs {runs}")
        else:
            self.emit("G6", "PASS", "no 3 consecutive text-only / non-visual pages")

    # ---- G7 visual anchor (by page_types) ----
    def g7_anchor(self, n, vis, spec):
        pts = spec.get("page_types") if spec else None
        if not pts or len(pts) != n:
            self.emit("G7", "WARN", "spec.page_types missing or length mismatch; skipping visual-anchor check")
            return
        miss = []
        for i, t in enumerate(pts):
            if t in VISUAL_PAGE_TYPES and not vis[i]:
                miss.append((i + 1, t))
        if miss:
            self.emit("G7", "FAIL", f"these pages need a visual anchor but lack one (page,type): {miss}")
        else:
            self.emit("G7", "PASS", "all content/data/table pages have a visual anchor")

    # ---- G8 placeholders ----
    def g8_placeholder(self, slides):
        hits = []
        for i, s in enumerate(slides, 1):
            txt = slide_text(s) + "\n" + notes_text(s)
            for pat in PLACEHOLDER_PATTERNS:
                for m in re.finditer(re.escape(pat), txt, flags=re.IGNORECASE):
                    frag = txt[max(0, m.start() - 15):m.start() + len(pat) + 15].replace("\n", " ")
                    hits.append((i, pat, frag))
        if hits:
            self.emit("G8", "FAIL", f"placeholders found: {hits}")
        else:
            self.emit("G8", "PASS", "no leftover placeholders")

    # ---- G9 banned template phrases ----
    def g9_forbidden(self, slides):
        hits = []
        for i, s in enumerate(slides, 1):
            txt = slide_text(s) + "\n" + notes_text(s)
            for pat in FORBIDDEN_PATTERNS:
                for m in re.finditer(pat, txt, flags=re.IGNORECASE):
                    ctx = txt[max(0, m.start() - 18):m.end() + 18].replace("\n", " ")
                    hits.append((i, m.group(0), ctx))
        if hits:
            self.emit("G9", "FAIL", f"banned template phrases found (page, phrase, context): {hits}")
        else:
            self.emit("G9", "PASS", "no banned template phrases")

    # ---- G10 image provenance hash check ----
    def g10_images(self, slides, figures, spec, mode):
        if figures is None:
            self.emit("G10", "WARN", "no --figures provided; cannot verify image provenance hashes")
            return
        manifest = {}
        for f in figures.get("figures", []):
            sha = (f.get("sha256") or "").lower()
            if sha:
                manifest[sha] = f
        docs_dir = spec.get("docs_dir") if spec else None

        # ledger self-consistency: every registered 'file' must exist and hash to the registered sha256
        # (prevents a manifest that only exists on paper, with no real image file behind it)
        ledger_bad = []
        for sha, f in manifest.items():
            fp = f.get("file") or ""
            tag = f.get("name") or sha[:8]
            if not fp or not os.path.isfile(fp):
                ledger_bad.append((tag, "file missing", fp or "-"))
                continue
            try:
                with open(fp, "rb") as fh:
                    real = hashlib.sha256(fh.read()).hexdigest()
            except Exception as e:
                ledger_bad.append((tag, f"read failed: {e}", fp))
                continue
            if real != sha:
                ledger_bad.append((tag, "sha256 mismatch vs ledger", fp))
        if ledger_bad:
            self.emit("G10", "FAIL", f"ledger inconsistent with files on disk (name, problem, file): {ledger_bad}")

        seen = {}
        unused_warn = set(manifest.keys())
        untrusted, wrong_source = [], []
        self.used_pdfcrop_labels = set()
        self.used_selfmade = 0
        for i, s in enumerate(slides, 1):
            for sh in walk_shapes(s.shapes):
                if not is_picture_shape(sh):
                    continue
                try:
                    blob = sh.image.blob
                except Exception:
                    continue
                sha = hashlib.sha256(blob).hexdigest()
                if sha in seen:
                    continue
                seen[sha] = i
                try:
                    wpx, hpx = sh.image.size
                except Exception:
                    wpx, hpx = -1, -1
                if sha not in manifest:
                    untrusted.append((i, f"{wpx}x{hpx}px", f"{len(blob)}B", sha[:16]))
                    continue
                unused_warn.discard(sha)
                entry = manifest[sha]
                st = entry.get("source_type")
                # record what is on the slides: paper figures (by label) / self-made count, for gate G15
                if st == PAPER_CROP and entry.get("label"):
                    self.used_pdfcrop_labels.add(str(entry["label"]))
                if st == SELF_MADE:
                    self.used_selfmade += 1
                if mode == "A":
                    origin = entry.get("origin") or ""
                    if st == PAPER_CROP:
                        if (not origin.lower().endswith(".pdf")
                                or not entry.get("page") or not entry.get("label")):
                            wrong_source.append((i, sha[:16], st, "pdf-crop but origin/page/label incomplete"))
                    elif st == "docs-asset":
                        if not docs_dir:
                            wrong_source.append((i, sha[:16], st, "missing spec.docs_dir, cannot verify"))
                        else:
                            try:
                                ok = os.path.abspath(origin).startswith(os.path.abspath(docs_dir))
                            except Exception:
                                ok = False
                            if not ok:
                                wrong_source.append((i, sha[:16], st, f"origin not inside docs_dir: {origin}"))
                    elif st == SELF_MADE:
                        rsn = str(entry.get("reason") or "").strip()
                        need = min_evidence_len(rsn)
                        if len(rsn) < need:
                            wrong_source.append((i, sha[:16], st,
                                                 f"self-made reason too short (<{need} chars): state why the material has no matching original figure"))
                    else:
                        wrong_source.append((i, sha[:16], st,
                                             "mode A allows only pdf-crop / docs-asset / self-made"))

        if untrusted:
            self.emit("G10", "FAIL",
                      f"untrusted provenance: images not registered in figures.json (page,size,bytes,sha16): {untrusted}")
        else:
            self.emit("G10", "PASS", f"all embedded images registered in the manifest ({len(seen)} after dedup)")
        # academic mode: unregistered images cannot prove their origin and count as non-compliant
        # (avoids a fake 'all image sources compliant' PASS)
        if mode == "A" and wrong_source:
            self.emit("G10", "FAIL", f"academic mode: non-compliant image sources (page,sha16,source_type,reason): {wrong_source}")
        elif mode == "A" and untrusted:
            self.emit("G10", "FAIL",
                      "academic mode with unregistered images: cannot prove they come from the paper; treated as untrusted provenance")
        elif mode == "A":
            self.emit("G10", "PASS",
                      "academic mode: all image sources compliant (pdf-crop / docs-asset inside docs_dir / documented self-made)")
        if unused_warn:
            self.emit("G10", "WARN",
                      f"registered in the manifest but unused in the deck (warning only): {sorted(unused_warn)}")

    # ---- G11 in-slide citation ----
    def g11_citation(self, slides, per_slide_shapes, figures, mode):
        if figures is None and mode != "A":
            self.emit("G11", "WARN", "no --figures provided; skipping in-slide citation check")
            return
        manifest = {}
        for f in (figures or {}).get("figures", []):
            sha = (f.get("sha256") or "").lower()
            if sha:
                manifest[sha] = f
        bad = []
        for i, (s, shs) in enumerate(zip(slides, per_slide_shapes), 1):
            has_pdf_crop = False
            has_any_pic = False
            for sh in shs:
                if is_picture_shape(sh):
                    has_any_pic = True
                    try:
                        sha = hashlib.sha256(sh.image.blob).hexdigest()
                    except Exception:
                        continue
                    e = manifest.get(sha)
                    if e and e.get("source_type") == "pdf-crop":
                        has_pdf_crop = True
                        break
            # academic mode: every image must be an original paper figure, so every page
            # containing any picture needs an in-slide citation
            if mode == "A":
                need = has_any_pic
            else:
                need = has_pdf_crop
            if not need:
                continue
            txt = slide_text(s)
            if not any(re.search(p, txt) for p in CITATION_PATTERNS):
                bad.append(i)
        if bad:
            self.emit("G11", "FAIL",
                      f"pages with figures but no in-slide citation: {bad} "
                      f"(paper figures must cite `Source: authors, Figure x`; self-made must be marked `Source: self-made`)")
        else:
            self.emit("G11", "PASS", "all pages with figures have an in-slide citation")

    # ---- G12 two-phase confirmation gate ----
    def g12_confirm(self, spec):
        if spec is None:
            self.emit("G12", "FAIL",
                      "missing spec: no outline_confirmed / style_confirmed; Stage 2 confirmation gate likely skipped")
            return
        problems = []
        for key, label in (("outline_confirmed", "outline"), ("style_confirmed", "style")):
            if spec.get(key) is not True:
                problems.append(f"{label} confirmation missing/false ({key})")

        # the evidence field accepts three spellings (SKILL.md §0.3 defines the shared confirm_evidence):
        #   1) "confirm_evidence": "User's exact words: ..."   <- documented form, proves both phases at once
        #   2) "confirm_evidence": {"outline": "...", "style": "..."}
        #   3) separate "outline_confirm_evidence" / "style_confirm_evidence" fields
        shared = spec.get("confirm_evidence")
        if isinstance(shared, dict):
            outline_ev, style_ev = str(shared.get("outline", "")), str(shared.get("style", ""))
        elif isinstance(shared, str):
            outline_ev = style_ev = shared.strip()
        else:
            outline_ev = style_ev = ""
        outline_ev = str(spec.get("outline_confirm_evidence") or outline_ev)
        style_ev = str(spec.get("style_confirm_evidence") or style_ev)
        for ev, label, field in ((outline_ev, "outline", "confirm_evidence/outline_confirm_evidence"),
                                 (style_ev, "style", "confirm_evidence/style_confirm_evidence")):
            need = min_evidence_len(ev)
            if len(ev) < need:
                problems.append(f"{label} confirmation evidence too short (<{need} chars, {field})")
        if problems:
            self.emit("G12", "FAIL", "Stage 2 confirmation gate skipped: " + "; ".join(problems))
        else:
            self.emit("G12", "PASS", "outline / style two-phase confirmation gate passed")

    # ---- G13 action title ----
    def g13_action_title(self, slides, per_slide_shapes, spec):
        pts = spec.get("page_types") if spec else None
        if not pts or len(pts) != len(slides):
            self.emit("G13", "WARN", "spec.page_types missing or length mismatch; skipping action-title gate")
            return
        bad = []
        for i, (s, shs, t) in enumerate(zip(slides, per_slide_shapes, pts), 1):
            if t not in TITLE_GATE_TYPES:
                continue
            best, best_pt = None, -1.0
            for sh in shs:
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue
                b = shape_bounds_in(sh)
                if b is None or b[1] >= 1.2:   # only shapes with top < 1.2in count as the title zone
                    continue
                fp = max_font_pt_of_shape(sh)
                if fp > best_pt:
                    best_pt, best = fp, sh.text_frame.text.strip()
            if best is None:
                bad.append((i, t, "no title text box"))
                continue
            if best in TITLE_BLACKLIST:
                bad.append((i, t, f"title is a section name in the blacklist: {best}"))
            else:
                min_len = 10 if cjk_char_count(best) >= 5 else 25
                if len(best) < min_len:
                    bad.append((i, t, f"title too short (<{min_len} chars): {best}"))
        if bad:
            self.emit("G13", "FAIL", f"action titles failing the gate (page,type,reason): {bad}")
        else:
            self.emit("G13", "PASS", "content/data/timeline/closing titles are all claim sentences")

    # ---- G14 fonts and default shadows (WARN-only) ----
    def g14_fonts_shadow(self, slides, per_slide_shapes):
        latin, ea = set(), set()
        shadow_inherit = 0
        for sh in (sh for shs in per_slide_shapes for sh in shs):
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    if para.font.name:
                        latin.add(para.font.name)
                    for r in para.runs:
                        if r.font.name:
                            latin.add(r.font.name)
                try:
                    for el in sh.text_frame._txBody.iter(qn("a:ea")):
                        tf = el.get("typeface")
                        if tf:
                            ea.add(tf)
                except Exception:
                    pass
            try:
                if sh.shadow.inherit is True:
                    shadow_inherit += 1
            except Exception:
                pass
        if len(latin) > 2:
            self.emit("G14", "WARN", f"more than 2 latin font families ({len(latin)}): {sorted(latin)}")
        else:
            self.emit("G14", "PASS", f"latin font families: {len(latin)} {sorted(latin)}")
        self.emit("G14", "WARN" if shadow_inherit else "PASS",
                  f"shapes using default shadow (inherit=True): {shadow_inherit} (source of the AI look; drive it to zero)")
        self.emit("G14", "PASS", f"East-Asian font families: {sorted(ea) if ea else '(no explicit <a:ea> set)'}")

    # ---- G15 paper-figure usage audit ("if the paper has it, use it") ----
    def g15_paper_figs(self, inventory, spec, mode):
        if mode != "A" and inventory is None:
            return  # non-academic mode and no paper inventory: gate not applicable
        if inventory is None:
            self.emit("G15", "WARN",
                      "no --paper-figures (inventory of the paper's figures) provided; cannot audit figure usage")
            return
        items = inventory.get("items") or []
        figs = [str(i.get("label")) for i in items if i.get("kind") == "Figure" and i.get("label")]
        if not figs:
            self.emit("G15", "WARN", "no Figure captions recognized in the paper inventory; skipping usage audit")
            return
        excused = (spec or {}).get("unused_figures_reason")
        excused = excused if isinstance(excused, dict) else {}
        used = [lb for lb in figs if lb in self.used_pdfcrop_labels]
        missing, weak = [], []
        for lb in figs:
            if lb in self.used_pdfcrop_labels:
                continue
            reason = str(excused.get(lb) or "").strip()
            if not reason:
                missing.append(lb)
            elif len(reason) < min_evidence_len(reason):
                weak.append((lb, reason))
        if missing:
            self.emit("G15", "FAIL",
                      f"the paper has these Figures but the deck neither uses them nor gives a reason "
                      f"(must go into spec.unused_figures_reason): {missing}")
        if weak:
            self.emit("G15", "FAIL", f"excuse for unused paper figure too short (effectively no reason): {weak}")
        # if the paper has figures, self-made figures must not replace them (paper decks):
        # self-made count must not exceed the number of paper figures actually used
        if mode == "A" and self.used_selfmade > len(used):
            self.emit("G15", "FAIL",
                      f"self-made figures {self.used_selfmade} > paper figures used {len(used)}: "
                      f"self-made figures appear to be replacing the originals")
        if not missing and not weak:
            self.emit("G15", "PASS",
                      f"paper has {len(figs)} Figures: {len(used)} used, "
                      f"the remaining {len(figs) - len(used)} all have documented reasons")
        if self.used_selfmade:
            self.emit("G15", "WARN",
                      f"this deck also contains {self.used_selfmade} self-made figure(s) (registered as source_type=self-made): "
                      f"self-made figures must supplement, never replace, the paper's originals")

    # ---- evidence block ----
    def evidence(self, n, vis, spec):
        deck_name = os.path.basename(self.args.deck)
        k = sum(1 for v in vis if v)
        ratio = (k / n * 100) if n else 0.0
        # collect all FAILs first, then append them to the evidence block (avoid mutating while iterating)
        fails = [ln for ln in self.lines if ln.startswith("[FAIL]")]
        self.lines.append("")
        self.lines.append("==== QA EVIDENCE (deck_qa.py) ====")
        self.lines.append(f"deck: {deck_name}")
        self.lines.append(
            f"slides: {n}   notes_range: CN {self.args.notes_min}-{self.args.notes_max} chars / "
            f"EN {self.args.notes_min_en}-{self.args.notes_max_en} words"
            f"   visual_pages: {k}/{n} = {ratio:.0f}%")
        self.lines.append(f"FAIL: {self.fail_count}   WARN: {self.warn_count}")
        for f in fails:
            self.lines.append(f)
        self.lines.append("==== END QA EVIDENCE ====")


# ---------------------------------------------------------------- entry point
def load_json(path, label, lines):
    if not path:
        lines.append(f"[WARN] {label} not provided (related checks degrade to WARN)")
        return None
    if not os.path.isfile(path):
        lines.append(f"[WARN] {label} file not found: {path}")
        return None
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except Exception as e:
        lines.append(f"[WARN] {label} failed to parse: {e}")
        return None


def main():
    ap = argparse.ArgumentParser(description="universal-ppt-skill single-entry QA gate (G1-G15)")
    ap.add_argument("deck", help="path to the .pptx deck to inspect")
    ap.add_argument("--spec", default=None,
                    help="deck-spec.json (pages / visual_ratio_min / confirm gates / page_types / docs_dir)")
    ap.add_argument("--figures", default=None, help="figures.json (figure provenance manifest)")
    ap.add_argument("--paper-figures", default=None,
                    help="paper-figures.json (inventory of the paper's own figures; gate G15 audits "
                         "'every original figure used or excused'; defaults to paper-figures.json next to --figures)")
    ap.add_argument("--out", default=None, help="write the full report to this file (QA evidence)")
    ap.add_argument("--notes-min", type=int, default=80, help="min speaker-notes length for CN decks (characters)")
    ap.add_argument("--notes-max", type=int, default=120, help="max speaker-notes length for CN decks (characters)")
    ap.add_argument("--notes-min-en", type=int, default=60, help="min speaker-notes length for EN decks (words)")
    ap.add_argument("--notes-max-en", type=int, default=100, help="max speaker-notes length for EN decks (words)")
    ap.add_argument("--mode", choices=list("ABCDEF"), default=None,
                    help="A academic / B business / C teaching / D everyday / E government / F competition")
    args = ap.parse_args()

    if not os.path.isfile(args.deck):
        sys.stderr.write(f"[FATAL] file not found: {args.deck}\n")
        raise SystemExit(2)

    qa = QA(args)
    spec = load_json(args.spec, "--spec", qa.lines)
    figures = load_json(args.figures, "--figures", qa.lines)
    inv_path = args.paper_figures
    if not inv_path and args.figures:
        guess = os.path.join(os.path.dirname(os.path.abspath(args.figures)), "paper-figures.json")
        if os.path.isfile(guess):
            inv_path = guess
    inventory = load_json(inv_path, "--paper-figures", qa.lines) if inv_path else None
    mode = args.mode or (spec.get("mode") if spec else None)

    try:
        prs = Presentation(args.deck)
    except Exception as e:
        sys.stderr.write(f"[FATAL] cannot open pptx ({args.deck}): {e}\n")
        raise SystemExit(2)

    qa.canvas = (prs.slide_width / EMU_PER_IN if prs.slide_width else CANVAS_W_IN_DEFAULT,
                 prs.slide_height / EMU_PER_IN if prs.slide_height else CANVAS_H_IN_DEFAULT)
    qa.canvas = (round(qa.canvas[0], 3), round(qa.canvas[1], 3))

    qa.lines.append(f"# deck_qa.py  deck={args.deck}  mode={mode or '(unspecified)'}  canvas={qa.canvas[0]}x{qa.canvas[1]}in")
    qa.run(prs, spec, figures, inventory, mode)

    text = "\n".join(qa.lines) + "\n"
    print(text, end="")
    if args.out:
        try:
            os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
            with open(args.out, "w", encoding="utf-8") as fh:
                fh.write(text)
        except Exception as e:
            sys.stderr.write(f"[WARN] cannot write --out: {e}\n")

    raise SystemExit(1 if qa.fail_count else 0)


if __name__ == "__main__":
    main()
