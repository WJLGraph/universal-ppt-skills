#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""deck_qa.py — universal-ppt-skill 单入口质检门禁（Stage 5 可执行门禁）

把 references/qa-gates.md 的散文规则变成机器能判死的检查项（G1-G15）。
任何 FAIL -> exit code = 1。

用法:
  python scripts/deck_qa.py <deck.pptx> [--spec work/deck-spec.json]
        [--figures work/figures.json] [--paper-figures work/paper-figures.json]
        [--out work/qa_output.txt]
        [--notes-min 80] [--notes-max 120] [--mode A|B|C|D|E|F]
"""

import argparse
import hashlib
import json
import math
import os
import re
import sys

# ---------------------------------------------------------------- 依赖检查
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn
except ImportError:
    sys.stderr.write(
        "[FATAL] 未安装 python-pptx，无法运行 deck_qa.py。\n"
        "        请先执行: pip install python-pptx\n"
    )
    raise SystemExit(2)

# ---------------------------------------------------------------- 常量
CANVAS_W_IN_DEFAULT = 13.333
CANVAS_H_IN_DEFAULT = 7.5
BOUND_TOL_IN = 0.03          # 形状越界容差（英寸）
GIANT_H_IN = 8.0             # 超大形状阈值
GIANT_W_IN = 13.5
EMU_PER_PT = 12700
EMU_PER_IN = 914400

VISUAL_PAGE_TYPES = {"content", "data", "table"}   # G7 需视觉锚点（timeline 由节点+连线形状构成，属豁免）
EXEMPT_PAGE_TYPES = {"title", "statement", "closing", "section", "timeline"}  # G7 豁免
TITLE_GATE_TYPES = {"content", "data", "timeline", "closing"}   # G13 适用范围

PLACEHOLDER_PATTERNS = [
    "XXX", "lorem", "TODO", "[insert", "待补充", "此处", "{{", "}}", "占位",
]
FORBIDDEN_PATTERNS = [
    r"一句话总结", r"最有价值的后续方向", r"值得注意的是", r"总的来说",
    r"综上所述", r"从某种意义上", r"提供了新的视角", r"具有重要意义",
    r"未来可以进一步探索", r"不是.{0,15}?而是", r"不只是.{0,15}?更是",
    r"众所周知", r"极大地", r"显著提升",
]
TITLE_BLACKLIST = {
    "研究背景", "背景", "方法", "方法介绍", "方案", "实验结果", "结果", "结论",
    "总结", "创新点", "不足", "参考文献", "致谢", "目录", "汇报大纲", "项目介绍",
    "实验设置",
}
CITATION_PATTERNS = [
    r"来源[:：]", r"Figure\s*\d+", r"Fig\.\s*\d+", r"Table\s*\d+", r"图\s*\d+",
    r"自绘", r"自制", r"作者绘制",
]

# 图片素材来源类型（G10）：论文 deck 优先用原文图，自绘图只允许作为补充且必须留痕+说明理由
PAPER_CROP = "pdf-crop"
SELF_MADE = "self-made"
MODE_A_SOURCE_TYPES = {PAPER_CROP, "docs-asset", SELF_MADE}

# G4 断行点：中文标点与英文标点本身即可断行，判定"不可断行 token"时必须先在此切开
BREAK_CHARS = "，。、；：？！）】》」』％…—·,.;:?!)]}"


# ---------------------------------------------------------------- 通用工具
def walk_shapes(shapes):
    """递归展开 group shape，产出叶子形状。"""
    for sh in shapes:
        try:
            is_group = sh.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:
            is_group = False
        if is_group:
            yield from walk_shapes(sh.shapes)
        else:
            yield sh


def is_picture_shape(sh):
    try:
        return sh.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE)
    except Exception:
        return False


def is_visual_shape(sh):
    """视觉锚点定义（G5/G7 核心）：PICTURE / 原生图表 / 表格，其余一律不算。"""
    if is_picture_shape(sh):
        return True
    if getattr(sh, "has_chart", False):
        return True
    if getattr(sh, "has_table", False):
        return True
    return False


def shape_bounds_in(sh):
    """返回 (left, top, right, bottom) 英寸，缺坐标返回 None。"""
    left, top = sh.left, sh.top
    w, h = sh.width, sh.height
    if left is None or top is None or w is None or h is None:
        return None
    return (left / EMU_PER_IN, top / EMU_PER_IN,
            (left + w) / EMU_PER_IN, (top + h) / EMU_PER_IN)


def slide_text(slide):
    parts = []
    for sh in walk_shapes(slide.shapes):
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False):
            for row in sh.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return "\n".join(p for p in parts if p)


def notes_text(slide):
    try:
        if not slide.has_notes_slide:
            return ""
        return slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        return ""


def char_width_em(ch):
    """中文/全角按 1.0 em，ASCII 按 0.5 em。"""
    return 1.0 if ord(ch) > 0x2E80 else 0.5


def resolve_font_pt(run, para):
    size = getattr(run.font, "size", None) or getattr(para.font, "size", None)
    return size.pt if size is not None else 18.0


def estimate_text_block(sh):
    """估算文本框。返回 (est_h_pt, max_line_w_pt, usable_w_pt, frame_h_pt, frame_w_pt)。"""
    tf = sh.text_frame
    frame_w_pt = sh.width / EMU_PER_PT
    frame_h_pt = sh.height / EMU_PER_PT
    ml = tf.margin_left if tf.margin_left is not None else 0
    mr = tf.margin_right if tf.margin_right is not None else 0
    mt = tf.margin_top if tf.margin_top is not None else 0
    mb = tf.margin_bottom if tf.margin_bottom is not None else 0
    usable_w_pt = max(frame_w_pt - (ml + mr) / EMU_PER_PT, 1.0)

    total_lines = 0
    max_line_w = 0.0
    max_font = 0.0
    for para in tf.paragraphs:
        runs = list(para.runs)
        if not runs:
            if not para.text:
                total_lines += 0          # 空段落暂不计行
                continue
            runs = None
        if runs is None:
            seg_pt = resolve_font_pt(para, para)
            w = sum(char_width_em(c) * seg_pt for c in para.text)
            max_font = max(max_font, seg_pt)
        else:
            w = 0.0
            for r in runs:
                seg_pt = resolve_font_pt(r, para)
                w += sum(char_width_em(c) * seg_pt for c in r.text)
                max_font = max(max_font, seg_pt)
        max_line_w = max(max_line_w, w)
        n_lines = max(1, int(math.ceil(w / usable_w_pt)))
        total_lines += n_lines
    if max_font <= 0:
        max_font = 18.0
    line_h = max_font * 1.35
    est_h_pt = total_lines * line_h + (mt + mb) / EMU_PER_PT
    return est_h_pt, max_line_w, usable_w_pt, frame_h_pt, frame_w_pt


def longest_unbreakable_token(sh):
    """返回 (token, 宽度pt)：最长且无法在词内断行的 token。
    中文可逐字断行，故只把含 ≥6 个 ASCII 字符的连续片段视为不可断行 token
    （如 langgraph、softmax(QKᵀ/√d)·V、路径/URL），这类才会真正冲出文本框。
    注意：中文标点（、，。；：等）本身就是合法断行点，必须先在标点处切开，
    否则 "系：Autoformer、FEDformer、PatchTST" 会被误判为单个超长 token。"""
    best_tok, best_w = "", 0.0
    splitter = re.compile(r"[\s" + re.escape(BREAK_CHARS) + r"]+")
    for para in sh.text_frame.paragraphs:
        pt = 0.0
        for r in para.runs:
            pt = max(pt, resolve_font_pt(r, para))
        if pt <= 0:
            pt = resolve_font_pt(para, para)
        for tok in splitter.split(para.text):
            ascii_n = sum(1 for c in tok if ord(c) < 128 and not c.isspace())
            if ascii_n < 6:
                continue
            w = sum(char_width_em(c) * pt for c in tok)
            if w > best_w:
                best_tok, best_w = tok, w
    return best_tok, best_w


def max_font_pt_of_shape(sh):
    best = 0.0
    if not sh.has_text_frame:
        return best
    for para in sh.text_frame.paragraphs:
        if para.font.size is not None:
            best = max(best, para.font.size.pt)
        for r in para.runs:
            if r.font.size is not None:
                best = max(best, r.font.size.pt)
    return best


# ---------------------------------------------------------------- 主流程
class QA:
    def __init__(self, args):
        self.args = args
        self.lines = []
        self.fail_count = 0
        self.warn_count = 0
        self.used_pdfcrop_labels = set()   # 上片的论文原图（Figure 号）
        self.used_selfmade = 0             # 上片的自绘图张数

    def emit(self, tag, level, msg):
        self.lines.append(f"[{level}] {tag} {msg}")
        if level == "FAIL":
            self.fail_count += 1
        elif level == "WARN":
            self.warn_count += 1

    def run(self, prs, spec, figures, inventory, mode):
        slides = list(prs.slides)
        n = len(slides)
        total_pages = spec.get("pages") if spec else None

        # 基础派生数据
        per_slide_shapes = [list(walk_shapes(s.shapes)) for s in slides]
        per_slide_visual = [any(is_visual_shape(sh) for sh in shs) for shs in per_slide_shapes]

        self.g1_pages(n, total_pages)
        self.g2_notes(slides)
        self.g3_bounds(slides, per_slide_shapes)
        self.g4_overflow(slides, per_slide_shapes)
        self.g5_visual_ratio(n, per_slide_visual, spec)
        self.g6_consecutive(n, per_slide_visual)
        self.g7_anchor(n, per_slide_visual, spec)
        self.g8_placeholder(slides)
        self.g9_forbidden(slides)
        self.g10_images(slides, figures, spec, mode)
        self.g11_citation(slides, per_slide_shapes, figures, mode)
        self.g12_confirm(spec)
        self.g13_action_title(slides, per_slide_shapes, spec)
        self.g14_fonts_shadow(slides, per_slide_shapes)
        self.g15_paper_figs(inventory, spec, mode)

        self.evidence(n, per_slide_visual, spec)

    # ---- G1 页数 ----
    def g1_pages(self, n, total_pages):
        if total_pages is None:
            self.emit("G1", "WARN", f"spec 未提供 pages，跳过页数核对（实际 {n} 页）")
        elif n == total_pages:
            self.emit("G1", "PASS", f"slides={n} 与 spec.pages 一致")
        else:
            self.emit("G1", "FAIL", f"slides={n} 与 spec.pages={total_pages} 不一致")

    # ---- G2 讲稿字数 ----
    def g2_notes(self, slides):
        lo, hi = self.args.notes_min, self.args.notes_max
        bad = []
        for i, s in enumerate(slides, 1):
            c = len(notes_text(s))
            self.emit("G2", "PASS", f"page {i:02d} notes_chars={c}")
            if c < lo or c > hi:
                bad.append((i, c))
        if bad:
            self.emit("G2", "FAIL", f"讲稿字数越界 [{lo},{hi}]：{bad}")
        else:
            self.emit("G2", "PASS", f"全部讲稿字数落在 [{lo},{hi}]")

    # ---- G3 形状越界 ----
    def g3_bounds(self, slides, per_slide_shapes):
        cw, ch = self.canvas
        bad = []
        giants = []
        for i, shs in enumerate(per_slide_shapes, 1):
            for k, sh in enumerate(shs):
                b = shape_bounds_in(sh)
                if b is None:
                    continue
                l, t, r, bo = b
                if (l < -BOUND_TOL_IN or t < -BOUND_TOL_IN or
                        r > cw + BOUND_TOL_IN or bo > ch + BOUND_TOL_IN):
                    bad.append((i, k, round(l, 2), round(t, 2), round(r, 2), round(bo, 2)))
                w_in = r - l
                h_in = bo - t
                if h_in > GIANT_H_IN or w_in > GIANT_W_IN:
                    giants.append((i, k, round(w_in, 2), round(h_in, 2)))
        if bad:
            self.emit("G3", "FAIL",
                      f"形状越界（画布 {cw}x{ch}in，容差 {BOUND_TOL_IN}in）: (page,idx,left,top,right,bottom)={bad}")
        else:
            self.emit("G3", "PASS", "无形状越界")
        if giants:
            self.emit("G3", "WARN",
                      f"GIANT 形状 (page,idx,w_in,h_in) 阈值 h>{GIANT_H_IN} 或 w>{GIANT_W_IN}: {giants}")
        else:
            self.emit("G3", "PASS", "无 GIANT 形状")

    # ---- G4 文本溢出估算 ----
    def g4_overflow(self, slides, per_slide_shapes):
        over_h, over_w, dense_para = [], [], []
        for i, shs in enumerate(per_slide_shapes, 1):
            for k, sh in enumerate(shs):
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue
                if sh.width is None or sh.height is None:
                    continue
                est_h, max_line_w, usable_w, frame_h, frame_w = estimate_text_block(sh)
                if est_h > frame_h * 1.15:
                    over_h.append((i, k, round(est_h, 1), round(frame_h, 1)))
                # 宽度只判「不可断行的长 token」（中文可任意断行，多行段落整段宽度会误报）
                tok, tok_w = longest_unbreakable_token(sh)
                if tok_w > usable_w * 1.02:
                    over_w.append((i, k, tok, round(tok_w, 1), round(usable_w, 1)))
                # 单段文字远宽于容器：正常换行后行数过多，提示拆分/加框
                if max_line_w > usable_w * 2.5:
                    dense_para.append((i, k, round(max_line_w, 1), round(usable_w, 1)))
        if over_h:
            self.emit("G4", "FAIL",
                      f"高度溢出 (page,idx,est_h_pt,frame_h_pt)：{over_h}")
        else:
            self.emit("G4", "PASS", "无高度溢出")
        if over_w:
            self.emit("G4", "FAIL",
                      f"不可断行长 token 超框 (page,idx,token,token_w_pt,usable_w_pt)：{over_w}")
        else:
            self.emit("G4", "PASS", "无不可断行长 token 超框")
        if dense_para:
            self.emit("G4", "WARN",
                      f"段落文字量远超容器宽（建议拆句或加宽）(page,idx,para_w_pt,usable_w_pt)：{dense_para}")
        else:
            self.emit("G4", "PASS", "无超长单段")

    # ---- G5 视觉页占比 ----
    def g5_visual_ratio(self, n, vis, spec):
        need = 0.6
        if spec and spec.get("visual_ratio_min") is not None:
            need = float(spec["visual_ratio_min"])
        k = sum(1 for v in vis if v)
        ratio = k / n if n else 0.0
        msg = f"visual_pages={k}/{n} = {ratio * 100:.1f}% (需 >= {need * 100:.0f}%)"
        if ratio >= need:
            self.emit("G5", "PASS", msg)
        else:
            self.emit("G5", "FAIL", msg + " —— 视觉页不足（图片/原图/原生图表/表格才算）")

    # ---- G6 连续非视觉页 ----
    def g6_consecutive(self, n, vis):
        runs = []
        i = 0
        while i < n:
            if not vis[i]:
                j = i
                while j < n and not vis[j]:
                    j += 1
                if j - i >= 3:
                    runs.append((i + 1, j))
                i = j
            else:
                i += 1
        if runs:
            self.emit("G6", "FAIL", f"连续非视觉页 >=3：最长区间页码 {runs}")
        else:
            self.emit("G6", "PASS", "无连续 3 页纯文字/非视觉页")

    # ---- G7 视觉锚点（按 page_types） ----
    def g7_anchor(self, n, vis, spec):
        pts = spec.get("page_types") if spec else None
        if not pts or len(pts) != n:
            self.emit("G7", "WARN", "spec.page_types 缺失或长度不等页数，跳过视觉锚点核对")
            return
        miss = []
        for i, t in enumerate(pts):
            if t in VISUAL_PAGE_TYPES and not vis[i]:
                miss.append((i + 1, t))
        if miss:
            self.emit("G7", "FAIL", f"这些页需视觉锚点但缺失 (page,type)：{miss}")
        else:
            self.emit("G7", "PASS", "content/data/table 页均有视觉锚点")

    # ---- G8 占位符 ----
    def g8_placeholder(self, slides):
        hits = []
        for i, s in enumerate(slides, 1):
            txt = slide_text(s) + "\n" + notes_text(s)
            for pat in PLACEHOLDER_PATTERNS:
                for m in re.finditer(re.escape(pat), txt, flags=re.IGNORECASE):
                    frag = txt[max(0, m.start() - 15):m.start() + len(pat) + 15].replace("\n", " ")
                    hits.append((i, pat, frag))
        if hits:
            self.emit("G8", "FAIL", f"发现占位符 {hits}")
        else:
            self.emit("G8", "PASS", "无占位符残留")

    # ---- G9 禁语扫描 ----
    def g9_forbidden(self, slides):
        hits = []
        for i, s in enumerate(slides, 1):
            txt = slide_text(s) + "\n" + notes_text(s)
            for pat in FORBIDDEN_PATTERNS:
                for m in re.finditer(pat, txt):
                    ctx = txt[max(0, m.start() - 18):m.end() + 18].replace("\n", " ")
                    hits.append((i, m.group(0), ctx))
        if hits:
            self.emit("G9", "FAIL", f"命中去模板化禁语 (page,短语,上下文)：{hits}")
        else:
            self.emit("G9", "PASS", "无中文模板化禁语")

    # ---- G10 图片溯源哈希核对 ----
    def g10_images(self, slides, figures, spec, mode):
        if figures is None:
            self.emit("G10", "WARN", "未提供 --figures，无法核对图片溯源哈希")
            return
        manifest = {}
        for f in figures.get("figures", []):
            sha = (f.get("sha256") or "").lower()
            if sha:
                manifest[sha] = f
        docs_dir = spec.get("docs_dir") if spec else None

        # 台账自洽校验：登记的 file 必须存在，且其 sha256 与登记值一致
        # （防"台账自述与落盘文件不符"——只写一行 json 却拿不出对应图片文件）
        ledger_bad = []
        for sha, f in manifest.items():
            fp = f.get("file") or ""
            tag = f.get("name") or sha[:8]
            if not fp or not os.path.isfile(fp):
                ledger_bad.append((tag, "文件不存在", fp or "-"))
                continue
            try:
                with open(fp, "rb") as fh:
                    real = hashlib.sha256(fh.read()).hexdigest()
            except Exception as e:
                ledger_bad.append((tag, f"读取失败：{e}", fp))
                continue
            if real != sha:
                ledger_bad.append((tag, "sha256 与登记值不一致", fp))
        if ledger_bad:
            self.emit("G10", "FAIL", f"台账与落盘文件不符 (name,问题,file)：{ledger_bad}")

        seen = {}
        unused_warn = set(manifest.keys())
        untrusted, wrong_source = [], []
        self.used_pdfcrop_labels = set()
        self.used_selfmade = 0
        for i, s in enumerate(slides, 1):
            for sh in walk_shapes(s.shapes):
                if not is_picture_shape(sh):
                    continue
                try:
                    blob = sh.image.blob
                except Exception:
                    continue
                sha = hashlib.sha256(blob).hexdigest()
                if sha in seen:
                    continue
                seen[sha] = i
                try:
                    wpx, hpx = sh.image.size
                except Exception:
                    wpx, hpx = -1, -1
                if sha not in manifest:
                    untrusted.append((i, f"{wpx}x{hpx}px", f"{len(blob)}B", sha[:16]))
                    continue
                unused_warn.discard(sha)
                entry = manifest[sha]
                st = entry.get("source_type")
                # 记录"上片的是什么"：论文原图（按 Figure 号）/ 自绘图张数，供 G15 判定
                if st == PAPER_CROP and entry.get("label"):
                    self.used_pdfcrop_labels.add(str(entry["label"]))
                if st == SELF_MADE:
                    self.used_selfmade += 1
                if mode == "A":
                    origin = entry.get("origin") or ""
                    if st == PAPER_CROP:
                        if (not origin.lower().endswith(".pdf")
                                or not entry.get("page") or not entry.get("label")):
                            wrong_source.append((i, sha[:16], st, "pdf-crop 但 origin/page/label 不完整"))
                    elif st == "docs-asset":
                        if not docs_dir:
                            wrong_source.append((i, sha[:16], st, "缺 spec.docs_dir 无法核对"))
                        else:
                            try:
                                ok = os.path.abspath(origin).startswith(os.path.abspath(docs_dir))
                            except Exception:
                                ok = False
                            if not ok:
                                wrong_source.append((i, sha[:16], st, f"origin 不在 docs_dir 内：{origin}"))
                    elif st == SELF_MADE:
                        rsn = str(entry.get("reason") or "").strip()
                        if len(rsn) < 8:
                            wrong_source.append((i, sha[:16], st,
                                                 "self-made 缺 reason：须写明素材中为何没有对应原图（≥8 字）"))
                    else:
                        wrong_source.append((i, sha[:16], st,
                                             "mode A 仅允许 pdf-crop / docs-asset / self-made"))

        if untrusted:
            self.emit("G10", "FAIL",
                      f"来源不可信：未经 figures.json 登记的图片 (page,size,bytes,sha16)：{untrusted}")
        else:
            self.emit("G10", "PASS", f"全部嵌入图已在 manifest 登记（{len(seen)} 张去重）")
        # 学术模式：未登记的图无从判定来源，同样视为不合规（避免"图源均合规"假 PASS）
        if mode == "A" and wrong_source:
            self.emit("G10", "FAIL", f"学术模式图源不合规 (page,sha16,source_type,reason)：{wrong_source}")
        elif mode == "A" and untrusted:
            self.emit("G10", "FAIL",
                      "学术模式且存在未登记图片：无法证明来自论文原文，视为来源不可信")
        elif mode == "A":
            self.emit("G10", "PASS",
                      "学术模式：图源均合规 (pdf-crop / docs_dir 内 docs-asset / 已留痕的 self-made)")
        if unused_warn:
            self.emit("G10", "WARN",
                      f"manifest 登记但 deck 未使用的图（仅告警）：{sorted(unused_warn)}")

    # ---- G11 in-slide citation ----
    def g11_citation(self, slides, per_slide_shapes, figures, mode):
        if figures is None and mode != "A":
            self.emit("G11", "WARN", "未提供 --figures，跳过 in-slide citation 核对")
            return
        manifest = {}
        for f in (figures or {}).get("figures", []):
            sha = (f.get("sha256") or "").lower()
            if sha:
                manifest[sha] = f
        bad = []
        for i, (s, shs) in enumerate(zip(slides, per_slide_shapes), 1):
            has_pdf_crop = False
            has_any_pic = False
            for sh in shs:
                if is_picture_shape(sh):
                    has_any_pic = True
                    try:
                        sha = hashlib.sha256(sh.image.blob).hexdigest()
                    except Exception:
                        continue
                    e = manifest.get(sha)
                    if e and e.get("source_type") == "pdf-crop":
                        has_pdf_crop = True
                        break
            # 学术模式：任何图片都必须是论文原图，因此任何含图页都必须有 in-slide citation
            if mode == "A":
                need = has_any_pic
            else:
                need = has_pdf_crop
            if not need:
                continue
            txt = slide_text(s)
            if not any(re.search(p, txt) for p in CITATION_PATTERNS):
                bad.append(i)
        if bad:
            self.emit("G11", "FAIL",
                      f"含图页缺 in-slide citation：页码 {bad}（论文图须标 `来源：作者 · Figure x`，自绘图须标 `来源：自绘`）")
        else:
            self.emit("G11", "PASS", "含图页均有 in-slide citation")

    # ---- G12 两阶段确认门 ----
    def g12_confirm(self, spec):
        if spec is None:
            self.emit("G12", "FAIL",
                      "缺少 spec：无 outline_confirmed / style_confirmed，疑似跳过 Stage 2 确认门")
            return
        problems = []
        for key, label in (("outline_confirmed", "大纲"), ("style_confirmed", "风格")):
            if spec.get(key) is not True:
                problems.append(f"{label}确认缺失/为假({key})")

        # 证据字段兼容三种写法（SKILL.md §0.3 约定的是共享 confirm_evidence）：
        #   1) "confirm_evidence": "用户原话：…"          ← 文档约定，一段话同时证明两阶段
        #   2) "confirm_evidence": {"outline": "…", "style": "…"}
        #   3) "outline_confirm_evidence" / "style_confirm_evidence" 分字段
        shared = spec.get("confirm_evidence")
        if isinstance(shared, dict):
            outline_ev, style_ev = str(shared.get("outline", "")), str(shared.get("style", ""))
        elif isinstance(shared, str):
            outline_ev = style_ev = shared.strip()
        else:
            outline_ev = style_ev = ""
        outline_ev = str(spec.get("outline_confirm_evidence") or outline_ev)
        style_ev = str(spec.get("style_confirm_evidence") or style_ev)
        for ev, label, field in ((outline_ev, "大纲", "confirm_evidence/outline_confirm_evidence"),
                                 (style_ev, "风格", "confirm_evidence/style_confirm_evidence")):
            if len(ev) < 8:
                problems.append(f"{label}确认原话证据不足（{field} < 8 字符）")
        if problems:
            self.emit("G12", "FAIL", "跳过 Stage 2 确认门：" + "；".join(problems))
        else:
            self.emit("G12", "PASS", "大纲 / 风格两阶段确认门均通过")

    # ---- G13 Action title ----
    def g13_action_title(self, slides, per_slide_shapes, spec):
        pts = spec.get("page_types") if spec else None
        if not pts or len(pts) != len(slides):
            self.emit("G13", "WARN", "spec.page_types 缺失或长度不等，跳过 action title 门禁")
            return
        bad = []
        for i, (s, shs, t) in enumerate(zip(slides, per_slide_shapes, pts), 1):
            if t not in TITLE_GATE_TYPES:
                continue
            best, best_pt = None, -1.0
            for sh in shs:
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue
                b = shape_bounds_in(sh)
                if b is None or b[1] >= 1.2:   # top < 1.2in 才算标题区
                    continue
                fp = max_font_pt_of_shape(sh)
                if fp > best_pt:
                    best_pt, best = fp, sh.text_frame.text.strip()
            if best is None:
                bad.append((i, t, "无标题文本框"))
                continue
            if best in TITLE_BLACKLIST:
                bad.append((i, t, f"标题为章节名黑名单：{best}"))
            elif len(best) < 10:
                bad.append((i, t, f"标题过短({len(best)}字)：{best}"))
        if bad:
            self.emit("G13", "FAIL", f"action title 不合格 (page,type,原因)：{bad}")
        else:
            self.emit("G13", "PASS", "content/data/timeline/closing 页标题均为合规结论句")

    # ---- G14 字体与默认阴影（仅 WARN 报告） ----
    def g14_fonts_shadow(self, slides, per_slide_shapes):
        latin, ea = set(), set()
        shadow_inherit = 0
        for sh in (sh for shs in per_slide_shapes for sh in shs):
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    if para.font.name:
                        latin.add(para.font.name)
                    for r in para.runs:
                        if r.font.name:
                            latin.add(r.font.name)
                try:
                    for el in sh.text_frame._txBody.iter(qn("a:ea")):
                        tf = el.get("typeface")
                        if tf:
                            ea.add(tf)
                except Exception:
                    pass
            try:
                if sh.shadow.inherit is True:
                    shadow_inherit += 1
            except Exception:
                pass
        if len(latin) > 2:
            self.emit("G14", "WARN", f"latin 字族 > 2 种（{len(latin)}）：{sorted(latin)}")
        else:
            self.emit("G14", "PASS", f"latin 字族 {len(latin)} 种：{sorted(latin)}")
        self.emit("G14", "WARN" if shadow_inherit else "PASS",
                  f"使用默认阴影(inherit=True)的形状数：{shadow_inherit}（AI 味来源，建议清零）")
        self.emit("G14", "PASS", f"东亚字族去重集合：{sorted(ea) if ea else '（未显式设置 <a:ea>）'}")

    # ---- G15 论文原图使用核查（"原文有图就必须用"） ----
    def g15_paper_figs(self, inventory, spec, mode):
        if mode != "A" and inventory is None:
            return  # 非学术模式且无论文清单：本条不适用
        if inventory is None:
            self.emit("G15", "WARN",
                      "未提供 --paper-figures（论文图片清单），无法核对\"原文有图是否都被用上\"")
            return
        items = inventory.get("items") or []
        figs = [str(i.get("label")) for i in items if i.get("kind") == "Figure" and i.get("label")]
        if not figs:
            self.emit("G15", "WARN", "论文清单里未识别到 Figure 题注，跳过使用核查")
            return
        excused = (spec or {}).get("unused_figures_reason")
        excused = excused if isinstance(excused, dict) else {}
        used = [lb for lb in figs if lb in self.used_pdfcrop_labels]
        missing, weak = [], []
        for lb in figs:
            if lb in self.used_pdfcrop_labels:
                continue
            reason = str(excused.get(lb) or "").strip()
            if not reason:
                missing.append(lb)
            elif len(reason) < 8:
                weak.append((lb, reason))
        if missing:
            self.emit("G15", "FAIL",
                      f"论文已有这些 Figure 但 deck 未使用且未给理由"
                      f"（须写进 spec.unused_figures_reason）：{missing}")
        if weak:
            self.emit("G15", "FAIL", f"未使用论文图的原因过短（<8 字，等于没说）：{weak}")
        # 原文有图就不许用自绘图顶替（论文 deck）：自绘张数不得超过论文原图上片数
        if mode == "A" and self.used_selfmade > len(used):
            self.emit("G15", "FAIL",
                      f"自绘图 {self.used_selfmade} 张 > 论文原图上片 {len(used)} 张："
                      f"疑似用自绘图顶替原文图")
        if not missing and not weak:
            self.emit("G15", "PASS",
                      f"论文 {len(figs)} 个 Figure：上片 {len(used)} 个，"
                      f"其余 {len(figs) - len(used)} 个均已说明理由")
        if self.used_selfmade:
            self.emit("G15", "WARN",
                      f"本 deck 另有 {self.used_selfmade} 张自绘图（已留痕 source_type=self-made）："
                      f"自绘图不得替代论文原图，只作补充")

    # ---- 证据块 ----
    def evidence(self, n, vis, spec):
        deck_name = os.path.basename(self.args.deck)
        k = sum(1 for v in vis if v)
        ratio = (k / n * 100) if n else 0.0
        # 先收集所有 FAIL，再统一追加到证据块（避免边遍历边修改）
        fails = [ln for ln in self.lines if ln.startswith("[FAIL]")]
        self.lines.append("")
        self.lines.append("==== QA EVIDENCE (deck_qa.py) ====")
        self.lines.append(f"deck: {deck_name}")
        self.lines.append(
            f"slides: {n}   notes_range: {self.args.notes_min}-{self.args.notes_max}"
            f"   visual_pages: {k}/{n} = {ratio:.0f}%")
        self.lines.append(f"FAIL: {self.fail_count}   WARN: {self.warn_count}")
        for f in fails:
            self.lines.append(f)
        self.lines.append("==== END QA EVIDENCE ====")


# ---------------------------------------------------------------- 入口
def load_json(path, label, lines):
    if not path:
        lines.append(f"[WARN] 未提供 {label}（{label} 相关检查将降级为 WARN）")
        return None
    if not os.path.isfile(path):
        lines.append(f"[WARN] {label} 文件不存在：{path}")
        return None
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except Exception as e:
        lines.append(f"[WARN] {label} 解析失败：{e}")
        return None


def main():
    ap = argparse.ArgumentParser(description="universal-ppt-skill 单入口质检门禁")
    ap.add_argument("deck", help="待质检的 .pptx 路径")
    ap.add_argument("--spec", default=None, help="deck-spec.json（页数/视觉页占比/确认门/page_types/docs_dir）")
    ap.add_argument("--figures", default=None, help="figures.json（图片溯源 manifest）")
    ap.add_argument("--paper-figures", default=None,
                    help="paper-figures.json（论文图片清单，G15 判\"原文有图有没有被漏用\"；"
                         "缺省时自动找 --figures 同目录的 paper-figures.json）")
    ap.add_argument("--out", default=None, help="完整输出写入的文件（QA 留痕）")
    ap.add_argument("--notes-min", type=int, default=80)
    ap.add_argument("--notes-max", type=int, default=120)
    ap.add_argument("--mode", choices=list("ABCDEF"), default=None, help="A 学术/B 商务/C 教学/D 日常/E 政务/F 竞赛")
    args = ap.parse_args()

    if not os.path.isfile(args.deck):
        sys.stderr.write(f"[FATAL] 文件不存在：{args.deck}\n")
        raise SystemExit(2)

    qa = QA(args)
    spec = load_json(args.spec, "--spec", qa.lines)
    figures = load_json(args.figures, "--figures", qa.lines)
    inv_path = args.paper_figures
    if not inv_path and args.figures:
        guess = os.path.join(os.path.dirname(os.path.abspath(args.figures)), "paper-figures.json")
        if os.path.isfile(guess):
            inv_path = guess
    inventory = load_json(inv_path, "--paper-figures", qa.lines) if inv_path else None
    mode = args.mode or (spec.get("mode") if spec else None)

    try:
        prs = Presentation(args.deck)
    except Exception as e:
        sys.stderr.write(f"[FATAL] 无法打开 pptx（{args.deck}）：{e}\n")
        raise SystemExit(2)

    qa.canvas = (prs.slide_width / EMU_PER_IN if prs.slide_width else CANVAS_W_IN_DEFAULT,
                 prs.slide_height / EMU_PER_IN if prs.slide_height else CANVAS_H_IN_DEFAULT)
    qa.canvas = (round(qa.canvas[0], 3), round(qa.canvas[1], 3))

    qa.lines.append(f"# deck_qa.py  deck={args.deck}  mode={mode or '(未指定)'}  canvas={qa.canvas[0]}x{qa.canvas[1]}in")
    qa.run(prs, spec, figures, inventory, mode)

    text = "\n".join(qa.lines) + "\n"
    print(text, end="")
    if args.out:
        try:
            os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
            with open(args.out, "w", encoding="utf-8") as fh:
                fh.write(text)
        except Exception as e:
            sys.stderr.write(f"[WARN] 无法写入 --out：{e}\n")

    raise SystemExit(1 if qa.fail_count else 0)


if __name__ == "__main__":
    main()
